#!/usr/bin/env python3
"""Render the deck to an HTML mock so the browser can do the text wrapping.

A character-count estimator is far too pessimistic for real font metrics.  Instead every shape
is placed absolutely at its true EMU geometry with real point sizes, pictures embedded as
base64, and the browser is asked where the text actually ended up.  measure() in the page then
compares each text box's content height against its declared height.
"""
import base64, html, io, pathlib
from pptx import Presentation
from pptx.util import Emu

SRC = '/Users/daddycl/Downloads/20261019-collabMeeting-v2.pptx'
OUT = pathlib.Path(__file__).parent / 'mock.html'
PX = 96.0            # px per inch
EMU_IN = 914400.0


def px(v):
    return (v or 0) / EMU_IN * PX


prs = Presentation(SRC)
SW, SH = px(prs.slide_width), px(prs.slide_height)
parts = []

for idx, slide in enumerate(prs.slides, 1):
    hidden = slide._element.get('show') == '0'
    body = []
    for sh in slide.shapes:
        l, t, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        style = (f'left:{l:.1f}px;top:{t:.1f}px;width:{w:.1f}px;height:{h:.1f}px;')
        img = None
        try:
            img = sh.image
        except Exception:
            img = None
        if img is not None:
            b64 = base64.b64encode(img.blob).decode()
            body.append(f'<div class="shape img" style="{style}">'
                        f'<img src="data:{img.content_type};base64,{b64}"></div>')
            continue
        if sh.has_table:
            rows = []
            for r in sh.table.rows:
                cells = ''.join(
                    f'<td>{html.escape(c.text)}</td>' for c in r.cells)
                rows.append(f'<tr>{cells}</tr>')
            body.append(f'<div class="shape tbl" data-n="{html.escape(sh.name)}" '
                        f'style="{style}"><table>{"".join(rows)}</table></div>')
            continue
        if sh.has_chart:
            body.append(f'<div class="shape chart" style="{style}">chart</div>')
            continue
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        paras = []
        for p in sh.text_frame.paragraphs:
            txt = ''.join(r.text for r in p.runs)
            if not txt.strip():
                paras.append('<p>&nbsp;</p>')
                continue
            f = p.runs[0].font
            size = f.size.pt if f.size else 18
            bold = 'font-weight:700;' if f.bold else ''
            ital = 'font-style:italic;' if f.italic else ''
            col = ''
            try:
                if f.color is not None and f.color.type is not None:
                    col = f'color:#{f.color.rgb};'
            except Exception:
                pass
            ind = f'margin-left:{p.level * 18}px;' if p.level else ''
            paras.append(f'<p style="font-size:{size}pt;{bold}{ital}{col}{ind}">'
                         f'{html.escape(txt)}</p>')
        body.append(f'<div class="shape txt" data-n="{html.escape(sh.name)}" '
                    f'style="{style}">{"".join(paras)}</div>')
    tag = 'SUPP' if hidden else 'core'
    parts.append(f'<div class="wrap"><div class="lbl">{idx} · {tag}</div>'
                 f'<div class="slide" data-i="{idx}" '
                 f'style="width:{SW:.0f}px;height:{SH:.0f}px">{"".join(body)}</div></div>')

OUT.write_text(f"""<!doctype html><meta charset="utf-8"><title>deck mock</title>
<style>
 body{{background:#555;margin:0;padding:18px;font-family:Arial,Helvetica,sans-serif}}
 .wrap{{margin:0 auto 26px;width:{SW:.0f}px}}
 .lbl{{color:#fff;font:12px monospace;margin:2px 0}}
 .slide{{position:relative;background:#fff;overflow:visible;box-shadow:0 2px 10px #0006}}
 .shape{{position:absolute;box-sizing:border-box}}
 .txt p{{margin:0 0 .28em;line-height:1.22}}
 .img img{{width:100%;height:100%;object-fit:contain}}
 .tbl table{{width:100%;border-collapse:collapse;font-size:12pt}}
 .tbl td{{border:1px solid #ccc;padding:2px 5px}}
 .chart{{background:#eef;color:#669;font:11px monospace;padding:3px}}
 .bad{{outline:3px solid #e33;background:#e3333318}}
</style>
{''.join(parts)}
<script>
function measure(){{
  const out=[];
  document.querySelectorAll('.slide').forEach(s=>{{
    const si=s.dataset.i, sh=s.getBoundingClientRect();
    s.querySelectorAll('.txt,.tbl').forEach(b=>{{
      const r=b.getBoundingClientRect();
      const need=(b.tagName==='DIV'&&b.classList.contains('tbl'))
        ? b.querySelector('table').getBoundingClientRect().height
        : [...b.children].reduce((m,c)=>Math.max(m,c.offsetTop+c.offsetHeight),0);
      const over=need-r.height;
      const off=(r.bottom-sh.top)-sh.height;
      if(over>2||off>2){{b.classList.add('bad');
        out.push({{slide:+si,name:b.dataset.n,overflow:+over.toFixed(1),
                   pastSlideBottom:+off.toFixed(1)}});}}
    }});
  }});
  return out;
}}
window.__m=measure;
</script>""", encoding='utf-8')
print('wrote', OUT, f'({OUT.stat().st_size/1e6:.1f} MB)')
