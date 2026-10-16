#!/usr/bin/env python3
"""V3 of the collaboration deck: an introduction to the design pipeline.

House style for this version, set by the author:
  - titles are Arial 32, plain and descriptive; no commas, no punchline phrasing
  - no small grey footers anywhere
  - "--" instead of an em dash
  - main slides carry summary bullets, not full argument; detail lives in the hidden
    supplementary set

Built from the previous deck rather than the bare template, because python-pptx cannot copy
the embedded movie or the grouped LPS pathway artwork cleanly between files.
"""
import pathlib, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.packuri import PackURI

SRC = '/Users/daddycl/Downloads/202601019-collabMeeting.pptx'
OUT = '/Users/daddycl/Downloads/20261019-collabMeeting-v3.pptx'
REPO = pathlib.Path('/Users/daddycl/Desktop/Lab_computation_results.nosync/projects/'
                    'mmpbsa-macrocycle-benchmark')
FIG = REPO / 'figures'

NAVY = RGBColor(0x00, 0x2A, 0x5C)
INK = RGBColor(0x1F, 0x29, 0x33)
MID = RGBColor(0x44, 0x4F, 0x5A)
ACC = RGBColor(0xB5, 0x1F, 0x58)
GRN = RGBColor(0x2F, 0x6B, 0x3C)

BODY = (Inches(0.47), Inches(1.21), Inches(12.39), Inches(5.46))
TITLE_PT = 32

prs = Presentation(SRC)
BLANK = next(l for l in prs.slide_masters[0].slide_layouts if l.name == '1_Blank')
sldIdLst = prs.slides._sldIdLst


def clean(t):
    """Enforce the dash rule on anything written here."""
    return t.replace('—', '--').replace('–', '-')


def style(run, size=18, color=INK, bold=False, italic=False):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def set_title(slide, text, size=TITLE_PT):
    text = clean(text)
    assert ',' not in text, f'comma in title: {text!r}'
    ph = slide.shapes.title
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    style(r, size, NAVY, bold=True)
    return ph


def drop_extra_placeholders(slide):
    for sh in list(slide.placeholders):
        if sh.placeholder_format.idx != 0:
            sh._element.getparent().remove(sh._element)


_next = [0]


def new_slide():
    if not _next[0]:
        used = [int(re.search(r'slide(\d+)\.xml$', sl.part.partname).group(1))
                for sl in prs.slides]
        _next[0] = max(used) + 1
    s = prs.slides.add_slide(BLANK)
    s.part.partname = PackURI('/ppt/slides/slide%d.xml' % _next[0])
    _next[0] += 1
    drop_extra_placeholders(s)
    return s


def points(slide, items, top=1.10, left=0.47, width=12.39, size=16, gap=7):
    """Short summary bullets. Items are (level, text)."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.0))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (lvl, txt) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = clean(txt)
        style(r, size if lvl == 0 else size - 2, INK if lvl == 0 else MID, bold=(lvl == 0))
    return tb


def fit(img, box, top_pad=0.0, bottom_pad=0.0):
    from PIL import Image
    iw, ih = Image.open(img).size
    l, t, w, h = box
    t = t + Emu(int(top_pad * 914400))
    h = h - Emu(int((top_pad + bottom_pad) * 914400))
    sc = min(w / iw, h / ih)
    nw, nh = int(iw * sc), int(ih * sc)
    return int(l + (w - nw) / 2), int(t + (h - nh) / 2), nw, nh


def fig_slide(title, img, summary=(), size=TITLE_PT):
    s = new_slide()
    set_title(s, title, size)
    pad = 0.0
    if summary:
        points(s, [(0, t) for t in summary], top=1.06, size=15, gap=4)
        pad = 0.30 + 0.30 * len(summary)
    l, t, w, h = fit(str(img), BODY, top_pad=pad)
    s.shapes.add_picture(str(img), l, t, w, h)
    return s


def stat(slide, x, y, value, label, w=3.4, vsize=26, lsize=12, color=ACC):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = clean(value)
    style(r, vsize, color, bold=True)
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = clean(label)
    style(r2, lsize, MID)
    return tb


def table(slide, rows, left, top, width, height, fs=12, widths=None):
    gf = slide.shapes.add_table(len(rows), len(rows[0]),
                                Inches(left), Inches(top), Inches(width), Inches(height))
    tbl = gf.table
    if widths:
        for i, cw in enumerate(widths):
            tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ''
            p = cell.text_frame.paragraphs[0]
            txt = clean(str(val))
            bold = ri == 0 or txt.startswith('**')
            r = p.add_run(); r.text = txt.strip('*')
            style(r, fs, NAVY if ri == 0 else INK, bold=bold)
            if ci > 0:
                p.alignment = PP_ALIGN.RIGHT
    return gf


def section(title, sub=None):
    s = new_slide()
    t = s.shapes.title
    t._element.getparent().remove(t._element)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = clean(title)
    style(r, 40, NAVY, bold=True)
    if sub:
        p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = clean(sub)
        style(r2, 16, MID)
    return s


# ------------------------------------------------- keep / delete originals
orig = list(prs.slides)
KEEP = {1: 'title', 2: 'lps', 5: 'protocol', 8: 'reagents', 16: 'descriptors',
        19: 'mlchart', 20: 'fullsweep', 21: 'ablation', 22: 'anchors', 23: 'krasmetrics'}
kept = {name: orig[i - 1] for i, name in KEEP.items()}
ids = list(sldIdLst)
for i in range(len(orig), 0, -1):
    if i not in KEEP:
        prs.part.drop_rel(ids[i - 1].rId)
        sldIdLst.remove(ids[i - 1])


def strip_small_grey(slide):
    """Remove the tiny grey caption boxes this house style bans."""
    for sh in list(slide.shapes):
        if sh.is_placeholder or not sh.has_text_frame:
            continue
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        if not runs:
            continue
        sizes = [r.font.size.pt for r in runs if r.font.size]
        if sizes and max(sizes) <= 9:
            sh._element.getparent().remove(sh._element)


# ------------------------------------------------- edit the kept slides
t = kept['title']
for ph in t.placeholders:
    if ph.placeholder_format.idx == 1:
        tf = ph.text_frame; tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = 'Collaboration Meeting  |  19 October 2026'
        style(r, 20, INK)
        p2 = tf.add_paragraph(); r2 = p2.add_run()
        r2.text = ('Carlos Liu  |  PhD Student  |  Brown Lab  |  '
                   'Leslie Dan Faculty of Pharmacy  |  University of Toronto')
        style(r2, 18, MID)

pr_s = kept['protocol']
set_title(pr_s, 'MM/PBSA scoring protocol')
for sh in list(pr_s.shapes):
    if sh.has_text_frame and not sh.is_placeholder:
        sh._element.getparent().remove(sh._element)
drop_extra_placeholders(pr_s)
points(pr_s, [
    (0, 'One minimised structure per complex. No production trajectory'),
    (1, 'tleap parameterisation with in-house ncAA charges and parameters'),
    (1, 'OpenMM minimisation then MM/PBSA at interior dielectric 2'),
    (0, 'Seconds per compound'),
    (1, 'FEP+ costs GPU-days. This is roughly 1000x cheaper'),
    (1, 'Cheap enough to score a whole design panel'),
], top=1.30, left=0.47, width=5.6, size=16, gap=8)

ml = kept['mlchart']
set_title(ml, 'Machine-learned ranker')
strip_small_grey(ml)

an = kept['anchors']
set_title(an, 'Effect of anchor count')
drop_extra_placeholders(an)
strip_small_grey(an)

set_title(kept['descriptors'], 'Descriptor families tested')
for sh in kept['descriptors'].shapes:
    if sh.has_table:
        for row in sh.table.rows:
            c = row.cells[3]
            if '0.331' in c.text:
                c.text_frame.paragraphs[0].runs[0].text = '0.524-0.540 alone / 0.563 + Morgan'
set_title(kept['fullsweep'], 'Full method sweep')
set_title(kept['ablation'], 'Descriptor ablation')
set_title(kept['krasmetrics'], 'KRAS MD metrics against experiment')
strip_small_grey(kept['krasmetrics'])
for _sh in kept['krasmetrics'].shapes:
    if _sh.has_text_frame:
        for _p in _sh.text_frame.paragraphs:
            for _r in _p.runs:
                if 'Prehaps' in _r.text:
                    _r.text = _r.text.replace('Prehaps', 'Perhaps')

rg = kept['reagents']
set_title(rg, 'Earlier three-swap purchase')
_t = rg.shapes.title
_t.left, _t.top, _t.width = 0, 0, Inches(13.333)
for sh in list(rg.shapes):
    if sh.has_text_frame and 'Is MD validated' in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)


# ------------------------------------------------- new core slides
new = {}

new['pipeline'] = fig_slide('Design pipeline', FIG / 'fig1_pipeline.png', summary=[
    'Generative backbone design, sequence design, non-canonical substitution, '
    'MM/PBSA triage, then synthesis',
])

s = new_slide()
set_title(s, 'Benchmark design')
points(s, [
    (0, 'Two rules fixed before we looked at any result'),
    (1, 'Average within a receptor. Pooling across crystal forms rewards separating groups '
        'rather than ranking inside one'),
    (1, 'Compare compounds of the same size. Affinity tracks heavy-atom count and so does '
        'the score'),
    (0, 'The bar is a heavy-atom count rather than chance'),
    (1, 'Size alone reaches 0.487 Spearman and 0.602 pairwise accuracy'),
    (0, 'Reference set'),
    (1, '234 macrocycles across five targets with measured affinities and published FEP+ '
        'values'),
], top=1.35, size=17, gap=11)
new['design'] = s

new['pairwise'] = fig_slide(
    'Size-matched pairwise accuracy', FIG / 'size_matched_pairwise_accuracy.png', summary=[
        'Among peptides within three heavy atoms we pick the tighter binder 62.4% of the time. '
        'Chance is 50% and FEP+ reaches 83.9%',
    ])

new['parity'] = fig_slide(
    'Correlation with experiment by target', FIG / 'zha25_fepplus_parity_5families.png',
    summary=['Spearman +0.545 averaged over five targets against +0.831 for FEP+'])

new['confound'] = fig_slide(
    'Size dependence of the score', FIG / 'fig2_size_confound.png', summary=[
        'Experimental affinity tracks ligand size at -0.43 and the score tracks it at -0.55. '
        'This is why we evaluate at matched size',
    ])

new['sampling'] = fig_slide(
    'Ensemble sampling test', FIG / 'sampling_noise.png', summary=[
        'Four 10 ns replicas per compound. Replica noise is 6.81 kcal/mol against a '
        'between-compound signal of 5.74',
        'Longer sampling does not improve the ranking which is why we score one minimised '
        'structure',
    ])

new['rosetta'] = fig_slide(
    'Rosetta as a reference standard', FIG / 'rosetta_vs_experiment.png', summary=[
        'Rosetta interface energy reaches -0.149 against experiment and is wrong-signed on '
        'four of five targets',
        'It cannot be used to judge which protocol is better',
    ])

new['so3lr'] = fig_slide(
    'Machine-learned force field', FIG / 'so3lr_comparison.png', summary=[
        'SO3LR-SF reaches 0.292 which is below both MM/PBSA at 0.545 and a heavy-atom count '
        'at 0.487',
        'It has no solvation term. Desolvation is what carries most of our signal',
    ])

new['ff'] = fig_slide(
    'Force-field comparison', FIG / 'zha25_lpta_forcefield.png', summary=[
        'ff03 wins on the benchmark and ff19SB wins on the LptA designs',
        'We report this as a caution rather than a recommendation',
    ])

new['lpta212'] = fig_slide(
    'LptA designs in three branches', FIG / 'zha25_lpta_212.png', summary=[
        '212 designs: 20 canonical plus 96 swaps on those parents plus 96 swaps on a '
        'derivative-7 scaffold',
        'Scored within branch. Pooling the three branches produces an artefact',
    ])

new['lptacanon'] = fig_slide(
    'Canonical designs under size control', FIG / 'zha25_lpta_canonical.png', summary=[
        'Agreement with Rosetta is +0.671 raw and +0.221 once buried area is removed',
        'Minimisation depth makes no difference here',
    ])

new['swap'] = fig_slide(
    'Parent to swap comparison', FIG / 'zha25_lpta_parent_swap.png', summary=[
        '75 of 96 swaps score better than their parent and bury 70 square angstrom more '
        'interface',
        'Agreement on which swap helps falls from +0.209 to +0.053 under size control',
    ])

s = new_slide()
set_title(s, 'Derivative-7 building blocks purchased')
rows = [['reagent (ncAA code)', 'site', 'USD', 'rank /96', 'size-controlled'],
        ['Fmoc-L-Asp(OFm)-OH  (Q2B)', '3', '130.25', '1', '2'],
        ['Fmoc-L-Cys(STmp)-OH  (Q6W)', '3', '183.00', '3', '**67'],
        ['Fmoc-L-Dap(Ac)-OH  (QF2)', '14', '29.00', '6', '9'],
        ['Fmoc-L-3,3-diphenyl-Ala  (Q8S)', '14', '16.75', '8', '7'],
        ['Fmoc-L-Dab(Ac)-OH  (QEW)', '9', '103.25', '11', '5'],
        ['Fmoc-L-Phe(2,4-F2)-OH  (QC3)', '5', '64.60', '14', '14'],
        ['Fmoc-O-benzyl-L-Tyr  (Z1T)', '5', '20.00', '15', '18'],
        ['Fmoc-L-Phe(2-Me)-OH  (Q7I)', '5', '11.50', '16', '26'],
        ['Fmoc-L-Lys(propionyl)-OH  (Q8P)', '14', '43.00', '23', '11'],
        ['Fmoc-S-4-MeO-Bzl-L-Cys  (Q99)', '5', '2.50', '27', '**54']]
table(s, rows, left=0.45, top=1.30, width=8.55, height=4.1, fs=10.5,
      widths=[3.5, 0.75, 1.1, 1.3, 1.9])
stat(s, 9.30, 1.38, '$603.85', 'ten building blocks across four sites')
stat(s, 9.30, 2.70, '12 of 96', 'median predicted rank. A random pick would be 48')
stat(s, 9.30, 4.02, '8 of 10', 'hold a top-30 rank once buried area is removed', color=GRN)
points(s, [
    (0, 'Two reagents were ranked mainly by buried area. Fmoc-L-Cys(STmp)-OH at $183 falls '
        'from 3rd to 67th'),
], top=5.62, size=15)
new['purchase'] = s

s = new_slide()
set_title(s, 'Registered prospective test')
table(s, [['', 'registered before synthesis'],
          ['compounds', '10 de novo LptA designs. 2 canonical and 8 non-canonical'],
          ['size window', '149-153 heavy atoms so the set is size-matched'],
          ['predicted range', '-104.0 to -69.8 kcal/mol spanning the range'],
          ['ranking', 'fixed at selection and will not be revised']],
      left=0.6, top=1.38, width=7.4, height=1.9, fs=11.5, widths=[1.7, 5.7])
table(s, [['n', 'power at 0.30', '0.50', '0.75'], ['10', '0.12', '0.28', '0.67'],
          ['20', '0.22', '0.58', '0.96'], ['40', '0.43', '0.88', '1.00']],
      left=8.4, top=1.38, width=4.3, height=1.6, fs=10.5)
points(s, [
    (0, 'Ten compounds cannot confirm the ranking. Power is 12% at the measured effect size'),
    (0, 'Ten compounds can detect a large inversion and give us measured anchors'),
    (0, 'No LptA design has a measured affinity yet'),
], top=3.70, size=16, gap=12)
new['prospective'] = s


# ------------------------------------------------- supplementary
supp = [section('Supplementary',
                'Benchmark detail / rejected approaches / retired long-MD track / '
                'machine-learning detail / provenance')]

s = new_slide()
set_title(s, 'Per-target performance')
table(s, [['target', 'n', 'r', 'R2', 'rho', 'FEP+ rho', 'slope', 'offset'],
          ['Cyclin A/B', '21', '+0.609', '0.371', '+0.591', '+0.791', '0.202', '-59.4'],
          ['JAK2', '48', '+0.588', '0.345', '+0.595', '+0.788', '0.103', '-44.8'],
          ['KRAS', '45', '+0.300', '0.090', '+0.225', '+0.786', '0.125', '-74.8'],
          ['MCL-1', '57', '+0.693', '0.480', '+0.728', '+0.859', '0.099', '-69.0'],
          ['PCSK9', '63', '+0.551', '0.304', '+0.584', '+0.929', '0.208', '-64.9'],
          ['**macro', '**234', '**+0.548', '**0.318', '**+0.545', '**+0.831', '-', '-']],
      left=0.7, top=1.7, width=11.9, height=3.3, fs=13)
points(s, [(0, 'Computed within receptor group. Slope is the regression of score on '
               'experiment so its reciprocal gives the inflation factor')],
       top=5.35, size=14)
supp.append(s)

s = new_slide()
set_title(s, 'Pairwise accuracy with intervals')
table(s, [['target', 'identical size', '95% CI', 'within 3 atoms', '95% CI'],
          ['Cyclin A/B', '0.533', '[0.053, 1.000]', '0.674', '[0.454, 0.863]'],
          ['JAK2', '0.637', '[0.397, 0.831]', '0.598', '[0.437, 0.732]'],
          ['KRAS', '0.518', '[0.289, 0.727]', '0.496', '[0.339, 0.664]'],
          ['MCL-1', '0.549', '[0.328, 0.757]', '0.594', '[0.444, 0.733]'],
          ['PCSK9', '**0.765', '[0.613, 0.882]', '**0.757', '[0.671, 0.830]'],
          ['**macro MM/PBSA', '**0.601', '[0.472, 0.720]', '**0.624', '[0.555, 0.689]'],
          ['macro FEP+', '0.841', '[0.730, 0.912]', '0.839', '[0.782, 0.886]']],
      left=0.9, top=1.55, width=11.5, height=3.7, fs=12.5)
points(s, [(0, 'Chance is 0.500. Intervals resample compounds rather than pairs. '
               'Only PCSK9 is individually above chance')], top=5.5, size=14)
supp.append(s)

s = new_slide()
set_title(s, 'KRAS under two estimators')
table(s, [['KRAS', 'n', 'MM/PBSA rho', 'FEP+ rho'],
          ['within form 7YUZ', '17', '-0.307', '+0.576'],
          ['within form 8JJS', '28', '+0.046', '+0.782'],
          ['**macro over forms', '45', '**-0.130', '+0.679'],
          ['pooled over both', '45', '**+0.225', '+0.786']],
      left=1.7, top=1.6, width=9.9, height=2.0, fs=14)
points(s, [
    (0, 'The score is offset by -3.05 kcal/mol between the two crystal forms and the labels '
        'by -1.15 in the same direction'),
    (0, 'So the pooled +0.225 is credit for separating two groups rather than ranking within '
        'either. FEP+ is far less affected'),
    (0, 'This happens inside what is normally reported as a within-family number'),
], top=4.0, size=15, gap=10)
supp.append(s)

s = new_slide()
set_title(s, 'Single target against cohort numbers')
points(s, [
    (0, 'Both are real measurements of different things'),
    (1, '-0.101 is the KRAS same-20 panel. One target and twenty ligands'),
    (1, '+0.545 is the macro average over five targets and 234 compounds'),
    (0, 'KRAS is the hardest of the five for this method'),
    (1, 'Its own macro value over the two crystal forms is -0.130'),
    (1, 'Its size-matched accuracy of 0.496 is the only one of five sitting at chance'),
    (0, 'Four of five targets are positive and PCSK9 reaches 0.757'),
], top=1.5, size=17, gap=11)
supp.append(s)

s = new_slide()
set_title(s, 'The score ranks but is not a free energy')
stat(s, 0.8, 1.6, '0.099 to 0.208', 'per-family regression slope of score on experiment')
stat(s, 5.3, 1.6, '5 to 10 kcal/mol', 'of raw score per 1 kcal/mol of real affinity')
stat(s, 9.5, 1.6, '1.11 vs 0.85', 'macro RMSE after calibration against FEP+')
points(s, [
    (0, 'The raw number cannot be quoted as a binding free energy'),
    (0, 'A per-family linear calibration makes the RMSE interpretable but needs measured '
        'anchors in that family'),
    (0, 'Use it to order compounds within a series rather than to predict affinity'),
], top=3.5, size=17, gap=12)
supp.append(s)

s = new_slide()
set_title(s, 'How the intervals are computed')
points(s, [
    (0, 'Resample compounds rather than pairs'),
    (1, 'Pairs drawn from n compounds share compounds and are not independent'),
    (1, 'Pair-level gives [0.530, 0.693] where compound-level gives [0.472, 0.720]'),
    (0, 'Macro rather than pooled'),
    (1, 'Pairs are formed within a target then averaged across targets with equal weight'),
    (0, 'Exclusions applied throughout'),
    (1, 'One compound with an unrelieved clash dropped leaving 234'),
    (1, 'Censored labels dropped for pairwise analysis leaving 220'),
], top=1.5, size=16, gap=9)
supp.append(s)

supp.append(fig_slide('Consensus scoring test', FIG / 'consensus_fails.png', summary=[
    'Rank errors between the two methods correlate at +0.714 so the consensus at 0.445 loses '
    'to MM/PBSA alone at 0.545']))

supp.append(fig_slide('Sealed external test on a sixth target', FIG / 'external_12ca5.png',
                      summary=['17 peptidomimetics against 12CA5 with affinities sealed '
                               'before prediction. Every method landed below chance']))

s = new_slide()
set_title(s, 'Why the machine-learned force field underperforms')
points(s, [
    (0, 'It has no desolvation term. Four lines point the same way'),
    (1, 'Correlates +0.866 with our gas-phase term and only +0.610 with full MM/PBSA'),
    (1, 'Mean magnitude -131 kcal/mol sits in the gas-phase range rather than the solvated one'),
    (1, 'Wins only on MCL-1 which is the greasiest pocket and collapses on PCSK9 the most polar'),
    (1, 'The preprint states optimisation is performed in the gas phase'),
    (0, 'Adding PB desolvation to our gas-phase term moves macro rho from 0.394 to 0.545'),
    (1, 'That is the single largest lever in the comparison'),
], top=1.5, size=16, gap=9)
supp.append(s)

supp.append(fig_slide('Retired long-MD track', FIG / 'md_legacy_vs_current.png', summary=[
    'The 3 x 10 ns stability ranking inverts under 4 x 250 ns. The source repository marks '
    'the short-run files as historical context only']))

s = new_slide()
set_title(s, 'Two FEP+ baselines appear in this deck')
table(s, [['target', 'held-query splits', 'all 235 compounds'],
          ['Cyclin A/B', '0.792', '0.791'], ['JAK2', '0.792', '0.788'],
          ['KRAS', '**0.828', '**0.786'], ['MCL-1', '**0.903', '**0.859'],
          ['PCSK9', '**0.874', '**0.929'], ['**mean', '**0.838', '**0.831']],
      left=2.9, top=1.45, width=7.5, height=2.4, fs=11.5)
points(s, [
    (0, '0.838 is FEP+ on the held-out query compounds of each ten-anchor split'),
    (0, '0.831 is FEP+ over all 235 labelled compounds'),
    (0, 'Both reproduce exactly. Three of five targets differ so this is a cohort difference'),
    (0, 'The two should not be placed on one axis'),
], top=4.05, size=15, gap=8)
supp.append(s)

s = new_slide()
set_title(s, 'Where the numbers come from')
table(s, [['claim', 'source'],
          ['rho +0.545 and +0.831 per target', 'Table 3 / results/zha25_cohort_scores.csv'],
          ['pairwise 0.624 [0.555, 0.689]', 'Table 5 / size_matched_pairwise_accuracy.py'],
          ['KRAS pooled against macro', 'Table 4'],
          ['size confound -0.43 and -0.55', 'scripts/fig_size_confound.py'],
          ['replica noise 6.81 against 5.74', 'kras_all40_md_score_and_stability_metrics.csv'],
          ['Rosetta -0.149', 'results/rosetta_vs_experiment_fepplus.csv'],
          ['SO3LR-SF 0.292', 'so3lr_sf_evaluation/so3lr_sf_scores.csv'],
          ['consensus +0.714 and 0.445', 'so3lr_sf_evaluation/consensus.py'],
          ['12CA5 -0.891', 'results/zha25_12ca5_built_pose_scores.csv'],
          ['LptA branches and swaps', 'results/zha25_lpta_*.csv'],
          ['registered ten', 'results/prospective_10_denovo_size_matched.csv'],
          ['purchased reagent ranks', 'zha25_lpta_212_scores.csv / custom_registry_1083.tsv']],
      left=0.8, top=1.45, width=11.7, height=4.7, fs=11.5, widths=[4.4, 7.3])
supp.append(s)


# ------------------------------------------------- order and hide
ORDER = [kept['title'], kept['lps'], new['pipeline'], kept['protocol'], new['design'],
         new['pairwise'], new['parity'], new['confound'], new['sampling'], new['rosetta'],
         new['so3lr'], new['ff'], kept['mlchart'], kept['anchors'],
         new['lpta212'], new['lptacanon'], new['swap'], new['purchase'], kept['reagents'],
         new['prospective']]
CORE = len(ORDER)
ORDER += supp[:9] + [kept['descriptors'], kept['fullsweep'], kept['ablation']] \
         + supp[9:11] + [kept['krasmetrics']] + supp[11:]

by_part = {sl.part: sid for sid in list(sldIdLst)
           for sl in [prs.part.related_part(sid.rId)]}
for sid in list(sldIdLst):
    sldIdLst.remove(sid)
for sl in ORDER:
    sldIdLst.append(by_part[sl.part])

for sl in ORDER[:CORE]:
    if sl._element.get('show') == '0':
        del sl._element.attrib['show']
for sl in ORDER[CORE:]:
    sl._element.set('show', '0')

prs.save(OUT)
print(f'wrote {OUT}')
print(f'  {len(prs.slides)} slides: {CORE} core + {len(prs.slides) - CORE} supplementary')
